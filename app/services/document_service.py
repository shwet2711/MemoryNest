
from __future__ import annotations

import hashlib
import re
import uuid
from pathlib import Path

import fitz
from docx import Document as DocxDocument
from PIL import Image
from sqlalchemy.orm import Session

from app.db.chunk_model import DocumentChunk
from app.db.document_model import Document
from app.db.embedding_model import ChunkEmbedding
from app.db.event_model import DetectedEvent
from app.db.knowledge_graph_sync_model import KnowledgeGraphSyncJob
from app.services.ocr_service import (
    extract_image_text,
    extract_image_text_with_confidence,
)
from app.services.ocr_quality_service import assess_ocr_text
from app.services.pdf_service import extract_pdf_with_ocr
import logging

from app.core.config import NEO4J_ENABLED
from app.services.knowledge_graph_service import (
    delete_document_relationships,
)
from app.services.knowledge_graph_sync_service import (
    enqueue_delete_job,
)
from app.services.vector_store import delete_document_embeddings

logger = logging.getLogger(__name__)
PROJECT_ROOT = Path(
    __file__
).resolve().parents[2]

UPLOAD_DIR = (
    PROJECT_ROOT
    / "data"
    / "uploads"
)

EXTRACTED_DIR = (
    PROJECT_ROOT
    / "data"
    / "extracted"
)

CHUNK_OUTPUT_DIR = (
    PROJECT_ROOT
    / "data"
    / "chunks"
)

UPLOAD_DIR.mkdir(
    parents=True,
    exist_ok=True,
)

EXTRACTED_DIR.mkdir(
    parents=True,
    exist_ok=True,
)

CHUNK_OUTPUT_DIR.mkdir(
    parents=True,
    exist_ok=True,
)


ALLOWED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".jpg",
    ".jpeg",
    ".png",
}

MAX_FILE_SIZE = 20 * 1024 * 1024


def validate_file(
    filename: str,
    file_size: int,
) -> tuple[bool, str]:

    if not filename:
        return False, "Filename is missing."

    extension = Path(
        filename
    ).suffix.lower()

    if extension not in ALLOWED_EXTENSIONS:

        allowed = ", ".join(
            sorted(ALLOWED_EXTENSIONS)
        )

        return (
            False,
            f"Unsupported file type. Allowed: {allowed}",
        )

    if file_size <= 0:
        return False, "The uploaded file is empty."

    if file_size > MAX_FILE_SIZE:

        return (
            False,
            "File is too large. "
            "Maximum allowed size is 20 MB.",
        )

    return True, ""


def sanitize_filename(
    filename: str,
) -> str:

    original = Path(
        filename
    ).name

    cleaned = re.sub(
        r"[^a-zA-Z0-9._-]",
        "_",
        original,
    )

    return cleaned[:180]


def calculate_file_hash(
    file_path: Path,
) -> str:

    sha256 = hashlib.sha256()

    with file_path.open("rb") as file:

        while True:

            chunk = file.read(
                1024 * 1024
            )

            if not chunk:
                break

            sha256.update(chunk)

    return sha256.hexdigest()


def extract_pdf_text(
    file_path: Path,
) -> str:

    pages = extract_pdf_with_ocr(
        file_path
    )

    sections = []

    for page in pages:

        if page["text"].strip():

            sections.append(
                f"\n--- Page "
                f"{page['page_number']} ---\n"
            )

            sections.append(
                page["text"]
            )

    return "\n".join(
        sections
    ).strip()


def extract_docx_text(
    file_path: Path,
) -> str:

    document = DocxDocument(
        file_path
    )

    paragraphs = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    return "\n".join(
        paragraphs
    ).strip()


def extract_txt_text(
    file_path: Path,
) -> str:

    encodings = [
        "utf-8",
        "utf-8-sig",
        "cp1252",
        "latin-1",
    ]

    for encoding in encodings:

        try:

            return file_path.read_text(
                encoding=encoding
            ).strip()

        except UnicodeDecodeError:
            continue

    raise ValueError(
        "Unable to decode text file."
    )


def validate_image(
    file_path: Path,
) -> dict:

    with Image.open(
        file_path
    ) as image:

        return {
            "format": image.format,
            "width": image.width,
            "height": image.height,
            "mode": image.mode,
        }



def extract_text_with_metadata(
    file_path: Path,
) -> tuple[str, str, str, float | None]:
    """
    Extract text and return:
    (text, status, extraction_method, OCR confidence).

    OCR text is preserved even when its quality is low.
    Confidence is None for files that do not use OCR.
    """

    extension = file_path.suffix.lower()

    # PDF: combine page text and evaluate OCR pages.
    if extension == ".pdf":
        pages = extract_pdf_with_ocr(file_path)
        sections = []
        ocr_confidences = []
        ocr_methods = []

        for page in pages:
            page_text = page.get("text", "").strip()

            if not page_text:
                continue

            sections.append(
                f"\n--- Page {page['page_number']} ---\n"
            )
            sections.append(page_text)

            if page.get("method", "").startswith("ocr"):
                ocr_methods.append(page["method"])

                confidence = page.get("ocr_confidence")
                if confidence is not None:
                    ocr_confidences.append(float(confidence))

        text = "\n".join(sections).strip()

        if not text:
            return "", "ocr_required", "ocr", None

        if not ocr_methods:
            method = "text"
            confidence = None
        else:
            confidence = (
                round(
                    sum(ocr_confidences) / len(ocr_confidences),
                    2,
                )
                if ocr_confidences
                else None
            )

            if (
                "ocr_low_confidence" in ocr_methods
                or "ocr_failed" in ocr_methods
                or confidence is None
                or confidence < 60.0
            ):
                method = "ocr_low_confidence"
            else:
                method = "ocr"

        return text, "completed", method, confidence

    # DOCX: ordinary text extraction.
    if extension == ".docx":
        text = extract_docx_text(file_path)

        if text:
            return text, "completed", "text", None

        return "", "empty", "text", None

    # TXT: ordinary text extraction.
    if extension == ".txt":
        text = extract_txt_text(file_path)

        if text:
            return text, "completed", "text", None

        return "", "empty", "text", None

    # Images: capture OCR confidence directly.
    if extension in {".jpg", ".jpeg", ".png"}:
        validate_image(file_path)

        with Image.open(file_path) as image:
            text, confidence = extract_image_text_with_confidence(
                image
            )

        if not text:
            return "", "ocr_failed", "ocr", confidence

        quality = assess_ocr_text(text)

        if (
            confidence is None
            or confidence < 60.0
            or not quality.is_reliable
        ):
            method = "ocr_low_confidence"
        else:
            method = "ocr"

        return text, "completed", method, confidence

    raise ValueError(
        f"Unsupported file extension: {extension}"
    )


def extract_text(
    file_path: Path,
) -> tuple[str, str]:
    """
    Backward-compatible extraction interface.

    Existing callers continue receiving (text, status).
    """

    text, status, _, _ = extract_text_with_metadata(file_path)
    return text, status


def save_extracted_text(
    document_id: int,
    text: str,
) -> Path:

    output_path = (
        EXTRACTED_DIR
        / f"document_{document_id}.txt"
    )

    output_path.write_text(
        text,
        encoding="utf-8",
    )

    return output_path


def create_document(
    db: Session,
    user_id: int,
    filename: str,
    file_bytes: bytes,
) -> Document:

    file_size = len(
        file_bytes
    )

    valid, error = validate_file(
        filename,
        file_size,
    )

    if not valid:
        raise ValueError(error)

    safe_name = sanitize_filename(
        filename
    )

    extension = Path(
        safe_name
    ).suffix.lower()

    unique_id = uuid.uuid4().hex

    stored_filename = (
        f"{unique_id}{extension}"
    )

    file_path = (
        UPLOAD_DIR
        / stored_filename
    )

    file_path.write_bytes(
        file_bytes
    )

    try:

        document = Document(
            user_id=user_id,
            original_filename=filename,
            stored_filename=stored_filename,
            file_path=str(file_path),
            file_type=extension.lstrip(
                "."
            ),
            file_size=file_size,
            title=Path(
                filename
            ).stem,
            extraction_status="processing",
        )

        db.add(document)

        db.flush()

        (
            text,
            status,
            extraction_method,
            ocr_confidence,
        ) = extract_text_with_metadata(file_path)

        document.extraction_status = status
        document.extraction_method = extraction_method
        document.ocr_confidence = ocr_confidence
        if text:

            extracted_path = (
                save_extracted_text(
                    document.id,
                    text,
                )
            )

            document.extracted_text_path = (
                str(extracted_path)
            )

        db.commit()

        db.refresh(
            document
        )

        return document

    except Exception:

        db.rollback()

        if file_path.exists():
            file_path.unlink()

        raise


def get_user_documents(
    db: Session,
    user_id: int,
) -> list[Document]:

    return (
        db.query(Document)
        .filter(
            Document.user_id == user_id
        )
        .order_by(
            Document.created_at.desc()
        )
        .all()
    )


def get_document(
    db: Session,
    user_id: int,
    document_id: int,
) -> Document | None:

    return (
        db.query(Document)
        .filter(
            Document.id == document_id,
            Document.user_id == user_id,
        )
        .first()
    )



def delete_document(
    db: Session,
    user_id: int,
    document_id: int,
) -> bool:
    """
    Delete a user's document and persist Knowledge Graph cleanup.

    The graph cleanup job is committed with the relational deletion.
    If Neo4j is unavailable, the job remains pending for the worker.
    """

    document = get_document(
        db,
        user_id,
        document_id,
    )

    if document is None:
        return False

    # Queue graph cleanup before deleting the document.
    # If queueing fails, preserve the document and its related records.
    try:
        enqueue_delete_job(
            db,
            user_id=user_id,
            document_id=document_id,
        )
        db.flush()
    except Exception:
        db.rollback()
        logger.exception(
            "Could not queue graph cleanup for document %s, user %s",
            document_id,
            user_id,
        )
        raise

    chunks = (
        db.query(DocumentChunk)
        .filter(
            DocumentChunk.document_id == document_id,
            DocumentChunk.user_id == user_id,
        )
        .all()
    )

    for chunk in chunks:
        db.delete(chunk)

    embedding_records = (
        db.query(ChunkEmbedding)
        .filter(
            ChunkEmbedding.document_id == document_id,
            ChunkEmbedding.user_id == user_id,
        )
        .all()
    )

    for record in embedding_records:
        db.delete(record)

    events = (
        db.query(DetectedEvent)
        .filter(
            DetectedEvent.document_id == document_id,
            DetectedEvent.user_id == user_id,
        )
        .all()
    )

    for event in events:
        db.delete(event)

    db.delete(document)

    # Commit document deletion and the graph cleanup job together.
    db.commit()

    # Vector cleanup is independent of Neo4j.
    try:
        delete_document_embeddings(
            document_id=document_id,
            user_id=user_id,
        )
    except Exception:
        logger.exception(
            "Failed to delete embeddings for document %s, user %s",
            document_id,
            user_id,
        )

    # Try graph cleanup immediately. If it fails, retain the queued job.
    if NEO4J_ENABLED:
        try:
            delete_document_relationships(
                user_id=user_id,
                document_id=document_id,
            )

            delete_job = (
                db.query(KnowledgeGraphSyncJob)
                .filter(
                    KnowledgeGraphSyncJob.user_id == user_id,
                    KnowledgeGraphSyncJob.document_id == document_id,
                    KnowledgeGraphSyncJob.operation == "delete",
                )
                .first()
            )

            if delete_job is not None:
                delete_job.status = "completed"
                delete_job.last_error = None
                db.commit()

        except Exception:
            db.rollback()
            logger.exception(
                "Immediate graph cleanup failed for document %s, user %s. "
                "The queued cleanup job will be retried.",
                document_id,
                user_id,
            )

    # Remove local files only after the relational commit.
    paths_to_remove = [
        Path(document.file_path),
        CHUNK_OUTPUT_DIR / f"document_{document.id}_chunks.txt",
    ]

    if document.extracted_text_path:
        paths_to_remove.append(
            Path(document.extracted_text_path)
        )

    for path in paths_to_remove:
        try:
            if path.exists():
                path.unlink()
        except OSError:
            logger.exception(
                "Failed to remove file %s for deleted document %s",
                path,
                document_id,
            )

    return True
